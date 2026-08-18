"""Niche Research — Web Server.

FastAPI + uvicorn, bind 127.0.0.1:8780, chạy sau nginx basic auth.
Pattern giống PlannerY + Content Ultimate (nginx + systemd, không Docker).

Chạy local để test:  python3 server.py
Trên VPS:            systemctl start niche-research
"""
from __future__ import annotations

import json
import os
import re
import secrets
import shutil
import subprocess
import sys
import threading
import time
from datetime import datetime, timedelta
from pathlib import Path
from typing import Optional

import uvicorn
from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, StreamingResponse
from fastapi import Body
from fastapi.staticfiles import StaticFiles

# ── Đường dẫn ────────────────────────────────────────────────────────────────
# V3 (APPS.md app 3): MỌI dữ liệu (projects + data) trỏ NICHE_DATA_DIR
# (data/niche-research) qua env — code/web ở apps/, dữ liệu ở data/ đúng sổ địa
# bạ. Không đặt env (V2 standalone) → nằm cạnh code y như cũ.
HERE        = Path(__file__).resolve().parent
ORCH        = HERE / "orchestrator.py"
_DATA_GOC   = Path(os.environ.get("NICHE_DATA_DIR") or HERE)
PROJECTS    = _DATA_GOC / "projects"   # mỗi niche = 1 subfolder
WEB_DIR     = HERE / "web"
DATA_DIR    = "niche-data"             # phải khớp với orchestrator.py
ENV_PATH    = HERE / ".env"
DATA_PATH   = _DATA_GOC / "data"
USERS_FILE  = DATA_PATH / "users.json"
INVITES_FILE = DATA_PATH / "invites.json"
HTPASSWD    = HERE / ".htpasswd"
PROJECTS.mkdir(parents=True, exist_ok=True)
WEB_DIR.mkdir(exist_ok=True)
DATA_PATH.mkdir(parents=True, exist_ok=True)

import khoa_v3  # noqa: E402  (nguồn khóa V3 — KÉT OUTLIERY, xem docstring module)

# ── Settings keys ─────────────────────────────────────────────────────────────
SETTING_KEYS = [
    # (env_key, label, hint, is_secret)
    ("LLM_PROVIDER",       "LLM Provider",           "anthropic | glm | openai | grok | custom", False),
    ("ANTHROPIC_API_KEY",  "Anthropic API Key",       "sk-ant-...",                               True),
    ("GLM_API_KEY",        "GLM / Z.AI API Key",      "Z.AI key (model glm-5.2)",                 True),
    ("OPENAI_API_KEY",     "OpenAI API Key",          "sk-...",                                    True),
    ("GROK_API_KEY",       "Grok API Key",            "xai-...",                                   True),
    ("TRANSCRIPT_API_KEY", "Transcript API Key",       "transcriptapi.com — 1 credit/video",       True),
    ("GLM_THINKING",       "GLM Thinking",            "disabled (mặc định) | enabled",            False),
    ("SHORTS_GATE",        "Shorts Gate",             "on (mặc định, lọc Short) | off",           False),
    ("MAX_COMMENTS",       "Max Comments",            "Số comment tối đa / kênh (mặc định 300)",  False),
]

# ── .env helpers ─────────────────────────────────────────────────────────────
def _read_env() -> dict[str, str]:
    """Đọc .env → dict (bỏ qua comment và dòng trống)."""
    result: dict[str, str] = {}
    if not ENV_PATH.exists():
        return result
    for line in ENV_PATH.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, _, val = line.partition("=")
        result[key.strip()] = val.strip().strip("'\"")
    return result


def _merge_env(updates: dict[str, str]) -> None:
    """Ghi đè các key trong .env, giữ nguyên phần còn lại."""
    # Sanitize: không cho phép newline injection
    updates = {k: v.replace("\r", "").split("\n")[0] for k, v in updates.items()}

    lines = ENV_PATH.read_text(encoding="utf-8").splitlines() if ENV_PATH.exists() else []
    written: set[str] = set()
    out: list[str] = []

    for line in lines:
        stripped = line.strip()
        if stripped.startswith("#") or "=" not in stripped:
            out.append(line)
            continue
        key = stripped.partition("=")[0].strip()
        if key in updates:
            if key in written:
                out.append(f"# {line}")   # comment out duplicate
            else:
                val = updates[key]
                out.append(f"# {key}=" if val == "" else f"{key}={val}")
                written.add(key)
        else:
            out.append(line)

    for key, val in updates.items():
        if key not in written:
            out.append("" if not out or out[-1] != "" else "")
            out.append(f"{key}={val}")

    ENV_PATH.write_text("\n".join(out) + "\n", encoding="utf-8")


def _settings_payload() -> list[dict]:
    env = _read_env()
    result = []
    for key, label, hint, is_secret in SETTING_KEYS:
        val = env.get(key, "")
        result.append({
            "key":    key,
            "label":  label,
            "hint":   hint,
            "secret": is_secret,
            "set":    bool(val),
            "tail":   val[-4:] if is_secret and val else "",
            "value":  "" if is_secret else val,   # chỉ gửi giá trị rõ cho trường không-secret
        })
    return result


# ── SSO từ OUTLIERY (30/07/2026): MỘT hệ tài khoản duy nhất ──────────────────
# Bật bằng NICHE_TRUST_PROXY=1. OUTLIERY đã đăng nhập + phân quyền bộ phận × level, rồi tiêm
# X-Remote-User + X-Remote-Role (viewer/leader/owner, tính từ bảng hanh_dong của apps_registry)
# vào mỗi request đi qua proxy. Khi cờ bật, VAI LẤY TỪ HEADER — .env ADMIN_USERS + users.json
# không còn là nguồn sự thật (tránh cảnh mỗi app một bảng phân quyền lệch nhau).
# Van an toàn: chỉ tin khi client là LOOPBACK — app bind 127.0.0.1 nên chỉ proxy trên cùng máy
# tới được; lỡ sau này app mở ra LAN thì header giả từ ngoài vẫn bị chặn ở đây.
# 04/08/2026 thêm 'manager' (luật OUTLIERY: Manager KHÔNG BAO GIỜ ngang Owner):
# manager = toàn quyền VẬN HÀNH (mọi việc của leader + XÓA dự án/tài liệu) nhưng KHÔNG
# đụng quản trị (Settings/API key, users/invites — 'admin' giữ, chỉ Owner OUTLIERY nhận).
# V3 18/08 (Permissions v2 — DE.md mục 14): ƯU TIÊN X-Remote-Actions do gateway
# tính từ luật + tick lẻ + acting, chảy sang TỪNG request: quan_tri → admin ·
# toan_quyen → manager (vai_xoa) · tao → leader · còn lại seo (vai THẤP NHẤT của
# app — DEFAULT fail-closed). THIẾU hẳn header Actions (gateway đời cũ) →
# fallback X-Remote-Role DANH PHÁP MỚI (admin/manager/leader/viewer); vai lạ /
# user trắng → seo, KHÔNG BAO GIỜ trả None giữa chừng khi SSO đang bật (None →
# rơi về ADMIN_USERS = nguồn quyền cũ — đã bịt).


def _vai_tu_sso(request: Request) -> str | None:
    if os.environ.get("NICHE_TRUST_PROXY") != "1":
        return None
    if not (request.client and request.client.host in ("127.0.0.1", "::1", "localhost")):
        return None
    raw = request.headers.get("X-Remote-Actions")
    if raw is not None:
        hd = {s.strip() for s in raw.split(",") if s.strip()}
        if "quan_tri" in hd:
            return "admin"
        if "toan_quyen" in hd:
            return "manager"
        if "tao" in hd:
            return "leader"
        return "seo"
    vai = (request.headers.get("X-Remote-Role") or "").strip().lower()
    return {"admin": "admin", "manager": "manager", "leader": "leader"}.get(vai, "seo")


def _get_role(request: Request) -> str:
    """Trả 'admin' | 'manager' | 'leader' | 'seo'."""
    vai = _vai_tu_sso(request)
    if vai:
        return vai
    env = _read_env()
    admins_raw = env.get("ADMIN_USERS", "").strip()
    username = request.headers.get("X-Remote-User", "").strip()
    if not admins_raw:
        return "admin"                              # dev mode: mọi người là admin
    admins = {a.strip() for a in admins_raw.split(",") if a.strip()}
    if username in admins:
        return "admin"
    users = _read_users()
    return users.get(username, {}).get("role", "seo")


def _require_admin(request: Request) -> None:
    if _get_role(request) != "admin":
        raise HTTPException(403, "Chỉ Quản trị mới có quyền thực hiện thao tác này")


# LÀM GỌN V3 (luật Owner 16/08, khuôn RadarY/Content): khi SSO, MỌI cửa quản trị
# nội bộ (Settings/.env key · users/invites/vai) đóng 404 KỂ CẢ vai admin nội bộ
# — khóa nhập ở OUTLIERY General › API Keys, quyền cấp ở General › Permissions.
THONG_DIEP_QT = ("Quản trị đã chuyển về OUTLIERY — khóa API nhập ở General › "
                 "API Keys, quyền cấp ở General › Permissions")


def _chot_quan_tri(request: Request) -> None:
    """Chốt MỌI route quản trị: SSO bật → 404 vô điều kiện; V2 → admin như cũ."""
    if os.environ.get("NICHE_TRUST_PROXY") == "1":
        raise HTTPException(404, THONG_DIEP_QT)
    _require_admin(request)


def _require_manager(request: Request) -> None:
    """Vận hành NẶNG (xóa dự án/tài liệu) — manager trở lên; admin đương nhiên qua."""
    if _get_role(request) not in ("manager", "admin"):
        raise HTTPException(403, "Chỉ Manager/Quản trị mới có quyền thực hiện thao tác này")


def _require_leader(request: Request) -> None:
    if _get_role(request) not in ("leader", "manager", "admin"):
        raise HTTPException(403, "Chỉ Leader/Quản trị mới có quyền thực hiện thao tác này")


# ── Users / Invites helpers ───────────────────────────────────────────────────
_users_lock = threading.Lock()

def _read_users() -> dict:
    with _users_lock:
        if USERS_FILE.exists():
            try: return json.loads(USERS_FILE.read_text(encoding="utf-8"))
            except Exception: pass
    return {}

def _write_users(data: dict) -> None:
    with _users_lock:
        USERS_FILE.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")

def _read_invites() -> dict:
    if INVITES_FILE.exists():
        try: return json.loads(INVITES_FILE.read_text(encoding="utf-8"))
        except Exception: pass
    return {}

def _write_invites(data: dict) -> None:
    INVITES_FILE.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")

def _htpasswd_add(username: str, password: str) -> None:
    subprocess.run(
        ["htpasswd", "-bB", str(HTPASSWD), username, password],
        check=True, capture_output=True,
    )

def _htpasswd_del(username: str) -> None:
    subprocess.run(
        ["htpasswd", "-D", str(HTPASSWD), username],
        capture_output=True,
    )


# ── Process registry (in-memory) ──────────────────────────────────────────────
_procs: dict[str, subprocess.Popen] = {}
_lock  = threading.Lock()

# ── App ───────────────────────────────────────────────────────────────────────
app = FastAPI(title="Niche Research", docs_url=None, redoc_url=None)
app.mount("/web", StaticFiles(directory=str(WEB_DIR)), name="web")


@app.get("/", response_class=HTMLResponse)
def index():
    return (WEB_DIR / "index.html").read_text(encoding="utf-8")


@app.get("/api/health")
def health():
    """Hợp đồng V3 (apps.json health) — gateway/giám sát thăm dò sự sống."""
    return {"ok": True, "app": "niche-research"}


# ── Nguồn khóa V3 (KÉT OUTLIERY) — chuẩn bị khóa cho MỖI lần chạy pipeline ───
def _bom_khoa_youtube(d: Path, keys: list[str]) -> None:
    """Bơm khóa YouTube của KÉT vào competitors.txt của dự án — pipeline V2 đọc
    key AIza… từ CHÍNH file input (scripts/_common.load_input), V3 két bơm thay
    cho việc user dán tay. Dedup: key đã có trong file thì không ghi lại."""
    comp = d / "competitors.txt"
    txt = comp.read_text(encoding="utf-8") if comp.exists() else ""
    thieu = [k for k in keys if k not in txt]
    if thieu:
        with open(comp, "a", encoding="utf-8") as f:
            f.write("\n# OUTLIERY keys (khoa_v3 — tu bom moi run)\n"
                    + "\n".join(thieu) + "\n")


def _khoa_v3_chuan_bi(d: Path | None, *, youtube: bool = False,
                      llm: bool = False, transcript: bool = False,
                      du_phong: bool = False) -> dict[str, str]:
    """Lấy khóa từ KÉT cho một lần chạy orchestrator (mỗi run gọi MỚI, không
    cache — Owner đổi cấp phát là run sau ăn ngay). Trả env overlay cho _spawn.

    Cờ True = việc BẮT BUỘC cho run này → két thiếu là 503 thông điệp rõ, KHÔNG
    rơi về .env/Settings nội bộ (V3 không còn .env). du_phong = resume/watch
    không biết run cần stage LLM/transcript nào → việc nào két CÓ thì bơm sẵn,
    thiếu thì thôi (stage cần mà thiếu sẽ tự dừng bằng thông điệp pipeline)."""
    if not khoa_v3.bat():
        return {}
    try:
        cap = khoa_v3._goi_ket()
        env: dict[str, str] = {}
        if youtube:
            _bom_khoa_youtube(d, khoa_v3.khoa_theo_viec("quet_kenh", cap))
        if llm:
            env.update(khoa_v3.env_llm(cap))
        if transcript:
            env.update(khoa_v3.env_transcript(cap))
        if du_phong:
            for ham in (khoa_v3.env_llm, khoa_v3.env_transcript):
                try:
                    env.update(ham(cap))
                except RuntimeError:
                    pass
        return env
    except RuntimeError as e:
        raise HTTPException(503, str(e))


# ── API: me ──────────────────────────────────────────────────────────────────
@app.get("/api/me")
def me(request: Request):
    username = request.headers.get("X-Remote-User", "").strip() or "dev"
    # sso=True → giao diện GIẤU chip tài khoản + tab Thành viên: danh tính và vai đã hiện
    # ở topbar OUTLIERY, tài khoản quản ở một chỗ duy nhất bên đó.
    return {"username": username, "role": _get_role(request),
            "sso": _vai_tu_sso(request) is not None}


# ── API: projects ─────────────────────────────────────────────────────────────
@app.get("/api/projects")
def list_projects():
    out = []
    for d in sorted(PROJECTS.iterdir(), key=lambda p: p.stat().st_mtime, reverse=True):
        if not d.is_dir():
            continue
        nd      = d / DATA_DIR
        reports = sorted((d / "Report").glob("*.xlsx")) if (d / "Report").exists() else []
        summary = (d / "Report" / "SUMMARY.md").exists()

        state: dict = {}
        sf = nd / ".state.json"
        if sf.exists():
            try:
                state = json.loads(sf.read_text(encoding="utf-8"))
            except Exception:
                pass

        with _lock:
            running = d.name in _procs and _procs[d.name].poll() is None

        if running:
            status = "running"
        elif reports:
            status = "done"
        elif nd.exists():
            status = "paused"
        else:
            status = "new"

        wcfg = _load_json(nd / ".watch.json") or {}
        out.append({
            "name":          d.name,
            "status":        status,
            "mtime":         d.stat().st_mtime,
            "has_report":    bool(reports),
            "has_run":       nd.exists(),
            "doc_count":     len(reports) + (1 if summary else 0),
            "has_summary":   summary,
            "watch":         bool(wcfg.get("enabled")),
            "watchable":     (d / "competitors.txt").exists(),
            "llm":           state.get("llm", False),
            "skip_comments": state.get("skip_comments", False),
        })
    return out


@app.post("/api/projects")
async def create_project(request: Request):
    _require_leader(request)
    body = await request.json()
    name = (body.get("name", "") or "").strip()
    safe = re.sub(r"[^\w\-]", "_", name)[:60]
    if not safe:
        raise HTTPException(400, "Tên dự án không hợp lệ")
    d = PROJECTS / safe
    if d.exists():
        raise HTTPException(400, "Dự án đã tồn tại")
    d.mkdir()
    return {"ok": True, "name": safe}


@app.delete("/api/projects/{name}")
def delete_project(name: str, request: Request):
    _require_manager(request)   # 04/08: xóa dự án = VẬN HÀNH nặng — manager làm được
    d = _proj(name)
    with _lock:
        proc = _procs.pop(name, None)
    if proc and proc.poll() is None:
        proc.terminate()
        time.sleep(0.5)
    shutil.rmtree(d)
    return {"ok": True}


@app.get("/api/status/{name}")
def project_status(name: str):
    d = _proj(name)
    reports = sorted((d / "Report").glob("*.xlsx")) if (d / "Report").exists() else []
    with _lock:
        running = name in _procs and _procs[name].poll() is None
    return {
        "running":      running,
        "has_report":   bool(reports),
        "report_files": [r.name for r in reports],
        "has_summary":  (d / "Report" / "SUMMARY.md").exists(),
    }


# ── API: run / resume / llm ───────────────────────────────────────────────────
@app.post("/api/run")
async def run_project(request: Request,
    name:          str        = Form(...),
    competitors:   UploadFile = File(...),
    skip_comments: bool       = Form(False),
    force:         bool       = Form(False),
    deepdive:      bool       = Form(False),
    llm:           bool       = Form(False),
):
    _require_leader(request)
    safe = re.sub(r"[^\w\-]", "_", name.strip())[:60] or "project"
    d    = PROJECTS / safe
    d.mkdir(exist_ok=True)

    comp = d / "competitors.txt"
    comp.write_bytes(await competitors.read())

    # V3: khóa từ KÉT theo đúng stage run này cần (quet_kenh luôn; phan_tich khi
    # --llm; lay_transcript khi --deepdive) — thiếu là 503 rõ, không chạy mù.
    env_extra = _khoa_v3_chuan_bi(d, youtube=True, llm=llm, transcript=deepdive)

    argv = [sys.executable, str(ORCH), "run", str(comp), "--work", str(d)]
    if skip_comments: argv += ["--skip-comments"]
    if force:         argv += ["--force"]
    if deepdive:      argv += ["--deepdive"]
    if llm:           argv += ["--llm"]

    _spawn(safe, d, argv, env_extra)
    return {"name": safe, "status": "started"}


@app.post("/api/resume/{name}")
def resume_project(name: str, request: Request):
    _require_leader(request)
    d    = _proj(name)
    env_extra = _khoa_v3_chuan_bi(d, youtube=True, du_phong=True)
    argv = [sys.executable, str(ORCH), "resume", str(d)]
    _spawn(name, d, argv, env_extra)
    return {"name": name, "status": "resumed"}


@app.post("/api/llm/{name}/{agent}")
def run_llm(name: str, agent: str, request: Request):
    _require_leader(request)
    if agent not in {"namer", "auditor", "plan", "dna", "summary"}:
        raise HTTPException(400, f"Unknown agent: {agent}")
    d    = _proj(name)
    env_extra = _khoa_v3_chuan_bi(None, llm=True)
    key  = f"{name}:llm:{agent}"
    argv = [sys.executable, str(ORCH), "llm", str(d), "--agent", agent]
    _spawn(key, d, argv, env_extra)
    return {"name": name, "agent": agent, "status": "started"}


@app.post("/api/stop/{name}")
def stop_project(name: str, request: Request):
    _require_leader(request)
    with _lock:
        proc = _procs.get(name)
    if proc and proc.poll() is None:
        proc.terminate()
        return {"stopped": True}
    return {"stopped": False}


# ── API: settings (.env) ─────────────────────────────────────────────────────
@app.get("/api/settings")
def get_settings(request: Request):
    _chot_quan_tri(request)
    return {"settings": _settings_payload()}


@app.post("/api/settings")
async def save_settings(request: Request):
    _chot_quan_tri(request)
    body = await request.json()
    updates: dict[str, str] = body.get("updates", {})
    # Chỉ cho phép update các key đã khai báo
    allowed = {k for k, *_ in SETTING_KEYS}
    updates = {k: v for k, v in updates.items() if k in allowed}
    _merge_env(updates)
    return {"ok": True, "settings": _settings_payload()}


# ── API: users & invites ─────────────────────────────────────────────────────
@app.get("/api/users")
def list_users(request: Request):
    _chot_quan_tri(request)
    users   = _read_users()
    invites = _read_invites()
    # Lọc invite hết hạn trước khi trả
    now = datetime.utcnow().isoformat()
    valid_invites = {k: v for k, v in invites.items() if v.get("expires", "") > now}
    return {"users": users, "invites": valid_invites}


@app.post("/api/invite")
async def create_invite(request: Request):
    _chot_quan_tri(request)
    body = await request.json()
    role = body.get("role", "seo")
    if role not in {"leader", "seo"}:
        raise HTTPException(400, "role phải là leader hoặc seo")

    code    = secrets.token_hex(4).upper()   # 8 ký tự hex
    expires = (datetime.utcnow() + timedelta(days=7)).isoformat()
    invites = _read_invites()
    invites[code] = {"role": role, "expires": expires,
                     "created_by": request.headers.get("X-Remote-User","dev")}
    _write_invites(invites)
    return {"code": code, "role": role, "expires": expires}


@app.delete("/api/invite/{code}")
def revoke_invite(code: str, request: Request):
    _chot_quan_tri(request)
    invites = _read_invites()
    invites.pop(code, None)
    _write_invites(invites)
    return {"ok": True}


def _chan_dang_ky_khi_sso() -> None:
    # Một hệ tài khoản duy nhất: khi vai đến từ OUTLIERY thì đường đăng ký tài khoản riêng
    # (htpasswd cục bộ) phải ĐÓNG — tạo được cũng vô nghĩa và chỉ gây thêm một hệ tài khoản.
    if os.environ.get("NICHE_TRUST_PROXY") == "1":
        raise HTTPException(404, "Đăng ký đã tắt — tài khoản quản lý ở OUTLIERY (cổng 8000)")


@app.get("/api/register")
def register_page():
    """Trang đăng ký công khai — không cần auth (nginx để qua)."""
    _chan_dang_ky_khi_sso()
    return HTMLResponse("""<!doctype html><html lang="vi"><head>
<meta charset="utf-8"><title>Đăng ký — Niche Research</title>
<style>
body{font:14px/1.6 system-ui,sans-serif;background:#090c12;color:#e8edf4;
  display:flex;align-items:center;justify-content:center;min-height:100vh;margin:0}
.box{background:#121826;border:1px solid #243149;border-radius:10px;padding:32px 36px;width:340px}
h1{font-size:17px;margin-bottom:20px;color:#4c8fe0}
label{display:block;font-size:12.5px;color:#8b96a8;margin-bottom:4px}
input{width:100%;background:#182233;border:1px solid #243149;border-radius:7px;
  color:#e8edf4;padding:9px 12px;font:inherit;margin-bottom:14px;box-sizing:border-box}
button{width:100%;background:#4c8fe0;border:none;border-radius:8px;color:#0b1220;
  padding:10px;font:700 14px inherit;cursor:pointer}
button:hover{filter:brightness(1.08)}
.err{color:#D97C6C;font-size:13px;margin-bottom:10px}
.ok{color:#83A96F;font-size:13px;margin-bottom:10px}
</style></head><body>
<div class="box">
  <h1>🔭 Tham gia Niche Research</h1>
  <div id="msg"></div>
  <label>Mã mời</label><input id="code" placeholder="XXXXXXXX" />
  <label>Tên đăng nhập</label><input id="user" placeholder="username" autocomplete="off" />
  <label>Mật khẩu</label><input id="pass" type="password" placeholder="••••••••" />
  <button onclick="reg()">Tạo tài khoản</button>
</div>
<script>
async function reg(){
  const code=document.getElementById('code').value.trim().toUpperCase()
  const user=document.getElementById('user').value.trim()
  const pass=document.getElementById('pass').value
  const msg=document.getElementById('msg')
  if(!code||!user||!pass){msg.className='err';msg.textContent='Điền đầy đủ thông tin';return}
  const r=await fetch('/api/register',{method:'POST',headers:{'Content-Type':'application/json'},
    body:JSON.stringify({code,username:user,password:pass})})
  const d=await r.json()
  if(!r.ok){msg.className='err';msg.textContent=d.detail||'Lỗi'}
  else{msg.className='ok';msg.textContent='✓ Tài khoản đã tạo! Đang chuyển hướng…';
    setTimeout(()=>location.href='/',1500)}
}
</script></body></html>""")


@app.post("/api/register")
async def do_register(request: Request):
    """Đăng ký bằng mã mời — không cần auth."""
    _chan_dang_ky_khi_sso()
    body = await request.json()
    code     = (body.get("code", "") or "").strip().upper()
    username = (body.get("username", "") or "").strip()
    password = body.get("password", "") or ""

    if not re.match(r'^[a-zA-Z0-9_\-]{3,30}$', username):
        raise HTTPException(400, "Username 3–30 ký tự, chỉ chữ/số/_/-")
    if len(password) < 6:
        raise HTTPException(400, "Mật khẩu tối thiểu 6 ký tự")

    invites = _read_invites()
    inv = invites.get(code)
    if not inv:
        raise HTTPException(400, "Mã mời không hợp lệ hoặc đã hết hạn")
    if inv["expires"] < datetime.utcnow().isoformat():
        raise HTTPException(400, "Mã mời đã hết hạn")

    users = _read_users()
    if username in users:
        raise HTTPException(400, "Tên đăng nhập đã tồn tại")

    _htpasswd_add(username, password)
    users[username] = {"role": inv["role"], "created_at": datetime.utcnow().isoformat()}
    _write_users(users)

    invites.pop(code)
    _write_invites(invites)

    return {"ok": True, "username": username, "role": inv["role"]}


@app.delete("/api/users/{username}")
def remove_user(username: str, request: Request):
    _chot_quan_tri(request)
    caller = request.headers.get("X-Remote-User", "").strip()
    if username == caller:
        raise HTTPException(400, "Không thể xóa chính mình")
    _htpasswd_del(username)
    users = _read_users()
    users.pop(username, None)
    _write_users(users)
    return {"ok": True}


@app.post("/api/users/{username}/role")
async def set_role(username: str, request: Request):
    _chot_quan_tri(request)
    body = await request.json()
    role = body.get("role", "seo")
    if role not in {"leader", "seo"}:
        raise HTTPException(400, "role phải là leader hoặc seo")
    users = _read_users()
    if username not in users:
        raise HTTPException(404, "User không tồn tại")
    users[username]["role"] = role
    _write_users(users)
    return {"ok": True}


@app.post("/api/users/{username}/password")
async def set_password(username: str, request: Request):
    _chot_quan_tri(request)
    body = await request.json()
    password = body.get("password", "") or ""
    if len(password) < 6:
        raise HTTPException(400, "Mật khẩu tối thiểu 6 ký tự")
    _htpasswd_add(username, password)
    return {"ok": True}


# ── API: documents (thư viện tài liệu của dự án) ────────────────────────────
@app.get("/api/docs/{name}")
def list_docs(name: str):
    d  = _proj(name)
    rd = d / "Report"
    docs = []
    if rd.exists():
        for f in sorted(rd.glob("*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True):
            docs.append({
                "filename": f.name,
                "type":     "xlsx",
                "size":     f.stat().st_size,
                "mtime":    f.stat().st_mtime,
                "origin":   "pipeline" if f.name.endswith("_report.xlsx") else "upload",
            })
        sm = rd / "SUMMARY.md"
        if sm.exists():
            docs.append({
                "filename": "SUMMARY.md",
                "type":     "md",
                "size":     sm.stat().st_size,
                "mtime":    sm.stat().st_mtime,
                "origin":   "pipeline",
            })
    return {"docs": docs}


@app.post("/api/upload/{name}")
async def upload_report(name: str, file: UploadFile = File(...)):
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(400, "Chỉ chấp nhận file .xlsx")
    d = _proj(name)
    safe_file = re.sub(r"[^\w\-\. ]", "_", file.filename or "upload.xlsx")
    report_dir = d / "Report"
    report_dir.mkdir(exist_ok=True)
    (report_dir / safe_file).write_bytes(await file.read())
    return {"ok": True, "name": d.name, "filename": safe_file}


@app.delete("/api/doc/{name}/{filename}")
def delete_doc(name: str, filename: str, request: Request):
    _require_manager(request)   # 04/08: xóa tài liệu báo cáo = VẬN HÀNH nặng — manager làm được
    if ".." in filename or "/" in filename or "\\" in filename:
        raise HTTPException(400, "Invalid filename")
    d = _proj(name)
    f = d / "Report" / filename
    if not f.is_file():
        raise HTTPException(404, "File not found")
    f.unlink()
    return {"ok": True}


@app.get("/api/mdtext/{name}/{filename}")
def read_md(name: str, filename: str):
    if ".." in filename or "/" in filename or "\\" in filename or not filename.endswith(".md"):
        raise HTTPException(400, "Invalid filename")
    d = _proj(name)
    f = d / "Report" / filename
    if not f.is_file():
        raise HTTPException(404, "File not found")
    return {"content": f.read_text(encoding="utf-8", errors="replace")}


# ── API: view report (Excel → JSON) ──────────────────────────────────────────
@app.get("/api/report/{name}/{filename}")
def view_report(name: str, filename: str):
    if ".." in filename or "\\" in filename:
        raise HTTPException(400, "Invalid filename")
    d = _proj(name)
    f = d / "Report" / filename
    if not f.is_file():
        raise HTTPException(404, "File not found")
    try:
        import openpyxl
        wb = openpyxl.load_workbook(str(f), data_only=True, read_only=True)
        sheets = []
        for sh_name in wb.sheetnames:
            ws = wb[sh_name]
            rows = []
            total = 0
            for row in ws.iter_rows(values_only=True):
                total += 1
                if total <= 300:
                    rows.append([("" if c is None else str(c)) for c in row])
            # trim trailing empty rows
            while rows and all(c == "" for c in rows[-1]):
                rows.pop()
            sheets.append({"name": sh_name, "total_rows": total, "rows": rows})
        wb.close()
        return {"filename": filename, "sheets": sheets}
    except Exception as e:
        raise HTTPException(500, f"Không đọc được file: {e}")


# ── API: dashboard (Tầng 1 monitoring — portfolio tổng quan) ─────────────────
_gonogo_cache: dict[str, tuple[float, dict]] = {}

def _load_json(p: Path):
    try:
        return json.loads(p.read_text(encoding="utf-8"))
    except Exception:
        return None


def _parse_gonogo_xlsx(path: Path) -> dict:
    """Parse VERDICT / Attractiveness / Trend từ sheet Go-NoGo (label do 18_build_report.py ghi cố định).
    Cache theo mtime — file upload không đổi thì không parse lại."""
    try:
        mtime = path.stat().st_mtime
        hit = _gonogo_cache.get(str(path))
        if hit and hit[0] == mtime:
            return hit[1]
        import openpyxl
        wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
        out: dict = {}
        if "Go-NoGo" in wb.sheetnames:
            rows = list(wb["Go-NoGo"].iter_rows(max_row=40, values_only=True))
            for i, row in enumerate(rows):
                a = str(row[0]).strip() if row and row[0] else ""
                if a == "VERDICT" and len(row) > 1 and row[1]:
                    out["verdict"] = str(row[1])
                elif a.startswith("Attractiveness") and len(row) > 1 and row[1] is not None:
                    out["attractiveness"] = row[1]
                elif len(row) > 3 and row[3] == "Trend" and i + 1 < len(rows):
                    nxt = rows[i + 1]
                    if nxt and len(nxt) > 3 and nxt[3]:
                        out["trend"] = str(nxt[3])
        wb.close()
        _gonogo_cache[str(path)] = (mtime, out)
        return out
    except Exception:
        return {}


@app.get("/api/dashboard")
def dashboard():
    items = []
    for d in sorted(PROJECTS.iterdir(), key=lambda p: p.stat().st_mtime, reverse=True):
        if not d.is_dir():
            continue
        nd      = d / DATA_DIR
        reports = sorted((d / "Report").glob("*.xlsx")) if (d / "Report").exists() else []

        d1  = _load_json(nd / "decision1.json") or {}
        d2  = _load_json(nd / "decision2.json") or {}
        dem = _load_json(nd / "demand.json") or {}
        ba  = _load_json(nd / "bets_audited.json") or {}

        verdict = d1.get("decision")
        score   = d1.get("attractiveness")
        trend   = dem.get("trend")
        source  = "pipeline" if verdict else None

        if verdict is None and reports:
            # Dự án chỉ có xlsx (upload) → parse ngược sheet Go-NoGo
            pipe   = [r for r in reports if r.name.endswith("_report.xlsx")]
            gg     = _parse_gonogo_xlsx((pipe or reports)[-1])
            verdict = gg.get("verdict")
            score   = gg.get("attractiveness")
            trend   = trend or gg.get("trend")
            source  = "xlsx" if verdict else None

        beachhead = d2.get("decision") or ((d2.get("ranked") or [{}])[0].get("anchor"))

        bets = None
        if ba.get("final_bets"):
            bets = {"clone_now": 0, "clone": 0, "test": 0}
            for b in ba["final_bets"]:
                v = (b.get("verdict") or "").upper()
                if v == "CLONE NOW":
                    bets["clone_now"] += 1
                elif v == "CLONE":
                    bets["clone"] += 1
                elif v == "TEST":
                    bets["test"] += 1

        vj = nd / "videos.json"
        scanned_at = vj.stat().st_mtime if vj.exists() else None

        with _lock:
            running = d.name in _procs and _procs[d.name].poll() is None

        wcfg   = _load_json(nd / ".watch.json") or {}
        series = _series_of(d, limit=12)

        items.append({
            "name":           d.name,
            "running":        running,
            "verdict":        verdict,
            "attractiveness": score,
            "gate_reason":    d1.get("gate_reason"),
            "beachhead":      beachhead,
            "trend":          trend,
            "bets":           bets,
            "scanned_at":     scanned_at,
            "source":         source,
            "watch":          bool(wcfg.get("enabled")),
            "spark":          [s["excess"] for s in series],
            "doc_count":      len(reports) + (1 if (d / "Report" / "SUMMARY.md").exists() else 0),
            "mtime":          d.stat().st_mtime,
        })
    return {"projects": items}


# ── API: watch mode (Tầng 2) + signals (Tầng 3) + timeseries đồ thị ──────────
WATCH_FILE = ".watch.json"
OX_TH = 3.0

@app.get("/api/watch/{name}")
def get_watch(name: str):
    d   = _proj(name)
    cfg = _load_json(d / DATA_DIR / WATCH_FILE) or {}
    return {
        "enabled":       bool(cfg.get("enabled", False)),
        "interval_days": int(cfg.get("interval_days", 7)),
        "last_run":      cfg.get("last_run"),
        "watchable":     (d / "competitors.txt").exists(),
    }


@app.post("/api/watch/{name}")
async def set_watch(name: str, request: Request):
    _require_leader(request)
    d    = _proj(name)
    body = await request.json()
    enabled  = bool(body.get("enabled", False))
    interval = min(30, max(1, int(body.get("interval_days", 7))))
    if enabled and not (d / "competitors.txt").exists():
        raise HTTPException(400, "Dự án không có competitors.txt — không watch được (chỉ chứa báo cáo upload)")
    nd = d / DATA_DIR
    nd.mkdir(exist_ok=True)
    cfg = _load_json(nd / WATCH_FILE) or {}
    cfg["enabled"] = enabled
    cfg["interval_days"] = interval
    (nd / WATCH_FILE).write_text(json.dumps(cfg, ensure_ascii=False, indent=1), encoding="utf-8")
    return {"ok": True, **{k: cfg.get(k) for k in ("enabled", "interval_days", "last_run")}}


@app.post("/api/watch/{name}/run")
def run_watch_now(name: str, request: Request):
    """Quét ngay 1 watch cycle (không chờ scheduler)."""
    _require_leader(request)
    d = _proj(name)
    if not (d / "competitors.txt").exists():
        raise HTTPException(400, "Dự án không có competitors.txt — không watch được")
    env_extra = _khoa_v3_chuan_bi(d, youtube=True, du_phong=True)
    argv = [sys.executable, str(ORCH), "watch", str(d)]
    _spawn(name, d, argv, env_extra)
    return {"ok": True}


@app.get("/api/signals")
def all_signals():
    out = []
    for d in PROJECTS.iterdir():
        if not d.is_dir():
            continue
        sig = _load_json(d / DATA_DIR / "signals.json")
        for s in (sig or {}).get("signals", []):
            out.append({**s, "project": d.name})
    out.sort(key=lambda s: s.get("ts", 0), reverse=True)
    return {"signals": out[:100]}


def _series_of(d: Path, limit: int = 24) -> list[dict]:
    """snapshots/*.json → time-series gọn cho đồ thị (server tính, client chỉ vẽ)."""
    snap_dir = d / DATA_DIR / "snapshots"
    if not snap_dir.exists():
        return []
    series = []
    for f in sorted(snap_dir.glob("*.json"))[-limit:]:
        s = _load_json(f)
        if not s:
            continue
        vids = s.get("videos", {})
        outl = [v for v in vids.values() if v.get("ox", 0) >= OX_TH]
        series.append({
            "date":       s.get("date", f.stem),
            "ts":         s.get("ts"),
            "n_videos":   len(vids),
            "n_outliers": len(outl),
            "views":      sum(v.get("v", 0) for v in vids.values()),
            "excess":     sum(max(v.get("ex", 0), 0) for v in outl),
            "trend":      s.get("trend"),
        })
    return series


@app.get("/api/timeseries/{name}")
def timeseries(name: str):
    return {"series": _series_of(_proj(name))}


# ── Watch scheduler (thread nền — spawn watch cycle khi đến hạn) ─────────────
def _watch_scheduler():
    while True:
        try:
            for d in PROJECTS.iterdir():
                if not d.is_dir():
                    continue
                nd  = d / DATA_DIR
                cfg = _load_json(nd / WATCH_FILE)
                if not cfg or not cfg.get("enabled"):
                    continue
                if not (d / "competitors.txt").exists():
                    continue
                now      = time.time()
                last_run = cfg.get("last_run") or 0
                last_try = cfg.get("last_attempt") or 0
                interval = max(1, int(cfg.get("interval_days", 7))) * 86400
                # last_attempt chặn retry-loop đốt quota khi cycle chết giữa chừng (tối thiểu 6h)
                if now - last_run < interval or now - last_try < 6 * 3600:
                    continue
                cfg["last_attempt"] = now
                (nd / WATCH_FILE).write_text(json.dumps(cfg, ensure_ascii=False, indent=1), encoding="utf-8")
                try:
                    env_extra = _khoa_v3_chuan_bi(d, youtube=True, du_phong=True)
                except HTTPException:
                    continue   # két chưa cấp khóa → bỏ cycle này, không chạy mù
                argv = [sys.executable, str(ORCH), "watch", str(d)]
                _spawn(d.name, d, argv, env_extra)   # cùng key với run thủ công → không bao giờ chạy chồng
        except Exception:
            pass
        time.sleep(1800)   # kiểm mỗi 30 phút


@app.on_event("startup")
def _start_watch_scheduler():
    # startup event → chỉ chạy 1 lần trên app instance thực sự serve
    # (start thread ở module-level sẽ bị double khi `python server.py` import lại chính nó)
    # V3: NICHE_SCHEDULER=0 TẮT scheduler (khuôn RADARY_SCHEDULER — hệ thật C:\
    # vẫn tự watch theo lịch trên CÙNG dự án; V3 chạy song song là ĐỐT ĐÔI quota
    # + snapshot lệch; nghiệm thu watch bằng POST /api/watch/{name}/run tay).
    if os.environ.get("NICHE_SCHEDULER", "1") != "1":
        return
    threading.Thread(target=_watch_scheduler, daemon=True).start()


# ── API: log stream (SSE) ─────────────────────────────────────────────────────
@app.get("/api/log/{name}")
def stream_log(name: str):
    """Server-Sent Events: stream stdout.log của project."""
    log_path = PROJECTS / name / DATA_DIR / "stdout.log"

    def _iter():
        pos = 0
        # Gửi log hiện có
        if log_path.exists():
            with open(log_path, "r", encoding="utf-8", errors="replace") as f:
                for line in f:
                    clean = _strip_ansi(line.rstrip("\n"))
                    yield f"data: {json.dumps(clean)}\n\n"
            pos = log_path.stat().st_size

        # Tail log mới
        while True:
            with _lock:
                key     = name
                running = key in _procs and _procs[key].poll() is None

            if log_path.exists():
                size = log_path.stat().st_size
                if size > pos:
                    with open(log_path, "r", encoding="utf-8", errors="replace") as f:
                        f.seek(pos)
                        for line in f:
                            clean = _strip_ansi(line.rstrip("\n"))
                            yield f"data: {json.dumps(clean)}\n\n"
                    pos = log_path.stat().st_size

            if not running:
                yield "event: done\ndata: done\n\n"
                break

            time.sleep(0.35)

    return StreamingResponse(
        _iter(),
        media_type="text/event-stream",
        headers={
            "Cache-Control":    "no-cache",
            "X-Accel-Buffering": "no",   # tắt nginx buffer — bắt buộc để SSE hoạt động
        },
    )


# ── API: download ─────────────────────────────────────────────────────────────
@app.get("/api/download/{name}/{filename}")
def download(name: str, filename: str):
    if ".." in filename or "/" in filename or "\\" in filename:
        raise HTTPException(400, "Invalid filename")
    d = _proj(name)
    for base in [d / "Report", d]:
        f = base / filename
        if f.is_file():
            return FileResponse(
                str(f),
                filename=filename,
                media_type="application/octet-stream",
            )
    raise HTTPException(404, "File not found")


# ── Helpers ───────────────────────────────────────────────────────────────────
def _proj(name: str) -> Path:
    d = PROJECTS / re.sub(r"[^\w\-]", "_", name)
    if not d.is_dir():
        raise HTTPException(404, "Project not found")
    return d


def _strip_ansi(s: str) -> str:
    return re.sub(r"\033\[[0-9;]*m", "", s)


class _GiuCho:
    """Giữ chỗ trong _procs NGAY TRONG LOCK lúc kiểm tra — vá đua TOCTOU 19/08:
    kiểm-trong-lock nhưng spawn+đăng-ký ngoài lock nên 2 POST /api/run cùng lúc
    đều thấy 'chưa chạy' → LifeIn_ES bị chạy ĐÚP (2 orchestrator cùng giây, đốt
    đôi quota). poll() None để status coi là đang chạy trong cửa sổ spawn."""

    def poll(self):
        return None

    def terminate(self):
        pass


def _spawn(key: str, proj_dir: Path, argv: list[str],
           env_extra: dict[str, str] | None = None) -> None:
    """Spawn orchestrator subprocess; relay stdout → niche-data/stdout.log.
    env_extra = khóa/cấu hình từ KÉT V3 (chỉ sống trong env tiến trình con —
    không ghi .env)."""
    with _lock:
        existing = _procs.get(key)
        if existing and existing.poll() is None:
            return   # đang chạy rồi (hoặc đang giữ chỗ), bỏ qua
        _procs[key] = _GiuCho()   # RESERVE ngay trong lock — chặn request song song

    try:
        nd = proj_dir / DATA_DIR
        nd.mkdir(exist_ok=True)
        log_file = nd / "stdout.log"

        env = {**os.environ, "PYTHONUNBUFFERED": "1", **(env_extra or {})}
        proc = subprocess.Popen(
            argv,
            cwd=str(HERE),
            env=env,
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
        )
    except Exception:
        with _lock:                     # spawn hỏng → trả chỗ, không kẹt 'running'
            _procs.pop(key, None)
        raise

    def _relay():
        with open(log_file, "ab") as lf:
            assert proc.stdout is not None
            for chunk in iter(lambda: proc.stdout.read(256), b""):
                lf.write(chunk)
                lf.flush()
        proc.stdout.close()
        proc.wait()

    threading.Thread(target=_relay, daemon=True).start()

    with _lock:
        _procs[key] = proc


# ── Entry point ───────────────────────────────────────────────────────────────
if __name__ == "__main__":
    host = os.environ.get("NICHE_HOST", "127.0.0.1")
    port = int(os.environ.get("NICHE_PORT", "8780"))
    print(f"Niche Research → http://{host}:{port}", flush=True)
    uvicorn.run("server:app", host=host, port=port, reload=False)
