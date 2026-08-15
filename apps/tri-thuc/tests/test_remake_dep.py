"""Test bản đẹp PDF (Ý 4 Đợt 3) — module remake_dep, toàn provider giả/mock."""

import os

os.environ["MOCK_MODE"] = "true"

import csv
import io
from pathlib import Path

import docx
import pytest
from fastapi.testclient import TestClient

from src.main import CATALOG_HEADER, app, ghi_catalog
from src.remake_dep import (GHI_CHU, dung_trang_html, lam_sach_html, tao_ban_dep)


def _co_weasyprint():
    try:
        import weasyprint  # noqa: F401
        return True
    except Exception:
        return False


# Baseline hệ cũ trên Windows: WeasyPrint cần GTK (libgobject-2.0-0) — thiếu là không
# render PDF được. Bản đẹp là TÍNH NĂNG PHỤ (lỗi chỉ mất PDF, không chặn nạp liệu).
CAN_WEASYPRINT = pytest.mark.skipif(
    not _co_weasyprint(),
    reason="Thiếu WeasyPrint/GTK trên Windows — baseline hệ cũ; bản đẹp là tính năng phụ")


@pytest.fixture(autouse=True)
def bat_remake(monkeypatch):
    """conftest tắt REMAKE_DEP cho test cũ — file này bật lại."""
    monkeypatch.setenv("REMAKE_DEP", "true")


class ProviderGia:
    def __init__(self, tra):
        self.tra = tra
        self.so_lan = 0

    def generate(self, system, de_bai):
        self.so_lan += 1
        return self.tra


def _docx_mau(tmp_path, ten="tai-lieu.docx"):
    f = tmp_path / ten
    d = docx.Document()
    d.add_paragraph("Quy tắc 1: mỗi IP chỉ nuôi một tài khoản.")
    d.add_paragraph("Quy tắc 2: ngày đầu chỉ xem, không thao tác.")
    d.save(str(f))
    return f


@CAN_WEASYPRINT
def test_mock_writer_van_ra_file_pdf(tmp_path):
    """Đường vui đủ chuỗi: đọc docx → writer → critic ĐẠT → WeasyPrint ra PDF thật."""
    w = ProviderGia("<h2>Quy tắc</h2><p><mark>Mỗi IP một tài khoản.</mark></p>"
                    "<p>Ngày đầu chỉ xem, không thao tác.</p>")
    c = ProviderGia("ĐẠT")
    pdf = tao_ban_dep(_docx_mau(tmp_path), "Quy tắc nuôi", "KD-1", writer=w, critic=c)
    assert pdf is not None and pdf.exists()
    assert pdf.name == "tai-lieu_ban-dep.pdf"          # nằm cạnh bản gốc
    assert pdf.read_bytes()[:4] == b"%PDF"             # đúng là PDF thật
    assert (w.so_lan, c.so_lan) == (1, 1)              # ĐẠT → không sửa lại


@CAN_WEASYPRINT
def test_critic_bao_loi_writer_sua_lai_mot_vong(tmp_path):
    w = ProviderGia("<p>bản trình bày</p>")
    c = ProviderGia("LỖI\nThiếu ý Quy tắc 2.")
    pdf = tao_ban_dep(_docx_mau(tmp_path), "t", "KD-1", writer=w, critic=c)
    assert pdf is not None and pdf.exists()
    assert w.so_lan == 2                               # sinh + sửa theo phản biện
    assert c.so_lan == 1


def test_chi_nhan_docx(tmp_path):
    f = tmp_path / "ghi-chu.txt"
    f.write_text("chữ thường", encoding="utf-8")
    w = ProviderGia("<p>x</p>")
    assert tao_ban_dep(f, "t", "KD-1", writer=w, critic=w) is None
    assert w.so_lan == 0                               # không gọi model vô ích


def test_cong_tac_tat(tmp_path, monkeypatch):
    monkeypatch.setenv("REMAKE_DEP", "false")
    w = ProviderGia("<p>x</p>")
    assert tao_ban_dep(_docx_mau(tmp_path), "t", "KD-1", writer=w, critic=w) is None
    assert w.so_lan == 0


def test_docx_hong_khong_nem_loi(tmp_path):
    f = tmp_path / "hong.docx"
    f.write_bytes(b"x")                                # không phải docx thật
    assert tao_ban_dep(f, "t", "KD-1") is None         # lỗi → None, KHÔNG chặn upload


def test_lam_sach_cat_fence_va_script():
    tho = "```html\n<h2>Mục</h2><script>alert(1)</script><p>a</p>\n```"
    assert lam_sach_html(tho) == "<h2>Mục</h2><p>a</p>"


def test_lam_sach_bo_vo_trang_neu_model_buong():
    assert lam_sach_html("<html><body><p>a</p></body></html>") == "<p>a</p>"


def test_dong_ghi_chu_bat_buoc_va_escape_tieu_de():
    trang = dung_trang_html("Tiêu đề <hack>", "KD-1", "<p>x</p>")
    assert GHI_CHU in trang
    assert "&lt;hack&gt;" in trang                     # tiêu đề được escape


# ---- bước 3: gắn vào /upload + cột "Bản đẹp" trong _catalog.csv ----

def _docx_bytes():
    d = docx.Document()
    d.add_paragraph("Nội dung quy trình thử.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _upload(ten_file, noi_dung):
    from claims_v2 import client_claims  # V2: /upload cần claims Manager+
    return client_claims(app, "sep", "Kinh doanh", 5).post("/upload", data={
        "title": "Tài liệu thử", "keywords": "", "owner": "", "version": "v1",
        "department": "Kinh doanh", "doc_type": "Quy trình",
        "effective_status": "Còn hiệu lực", "access_level": "Công khai nội bộ",
        "min_level": "1"}, files={"file": (ten_file, noi_dung)})


def _catalog():
    f = Path(os.environ["KHO_TAI_LIEU"]) / "_catalog.csv"
    return list(csv.reader(f.open(encoding="utf-8-sig")))


@CAN_WEASYPRINT
def test_upload_docx_sinh_pdf_va_ghi_cot_ban_dep():
    r = _upload("quy-trinh.docx", _docx_bytes())  # mock writer/critic + weasyprint thật
    assert r.status_code == 200
    dong = _catalog()
    assert dong[0] == CATALOG_HEADER                    # header có cột "Bản đẹp"
    assert dong[1][15].endswith("_ban-dep.pdf")         # cột 16 = tên file bản đẹp
    ngan = Path(os.environ["KHO_TAI_LIEU"]) / dong[1][12]
    assert (ngan / dong[1][15]).read_bytes()[:4] == b"%PDF"  # PDF nằm cạnh bản gốc


def test_upload_tat_cong_tac_cot_ban_dep_rong(monkeypatch):
    monkeypatch.setenv("REMAKE_DEP", "false")
    assert _upload("quy-trinh.docx", _docx_bytes()).status_code == 200
    assert _catalog()[1][15] == ""                      # không PDF, upload vẫn OK


# ---- bước 4: GET /tai-ban-dep/{doc_code} — RBAC tái dùng _duoc_xem, 404 lặng lẽ ----

# V2: claims từ gateway thay users.txt — bảng bộ phận×level giữ nguyên ý cũ.
from claims_v2 import client_claims

HO_SO = {"ql": ("Kinh doanh", 4), "nv": ("Kinh doanh", 2)}


def _users_file(tmp_path, monkeypatch):
    """V2: không còn USERS_FILE — giữ chữ ký để call-site cũ nguyên vẹn (no-op)."""


def _dang_nhap(ten):
    return client_claims(app, ten, *HO_SO[ten])


def _chuan_bi_ban_dep(doc_code="KD-2026-0099", ngan="05_Kinh-doanh"):
    """Đặt sẵn PDF + dòng catalog cho 1 mã (mặc định 0099 = Mật min4 trong kho mock)."""
    kho = Path(os.environ["KHO_TAI_LIEU"])
    (kho / ngan).mkdir(parents=True, exist_ok=True)
    (kho / ngan / "x_ban-dep.pdf").write_bytes(b"%PDF-fake")
    ghi_catalog(kho, [doc_code] + [""] * 11 + [ngan, "", "", "x_ban-dep.pdf"])


def test_tai_ban_dep_dung_quyen_thi_duoc(tmp_path, monkeypatch):
    _users_file(tmp_path, monkeypatch)
    _chuan_bi_ban_dep()
    r = _dang_nhap("ql").get("/tai-ban-dep/KD-2026-0099")   # Manager KD, min4 <= 4
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/pdf")
    assert r.content == b"%PDF-fake"


def test_tai_ban_dep_thieu_level_404_lang_le(tmp_path, monkeypatch):
    """Bị chặn quyền phải 404 (không phải 403) — GIỐNG HỆT ca không-có-file."""
    _users_file(tmp_path, monkeypatch)
    _chuan_bi_ban_dep()
    r = _dang_nhap("nv").get("/tai-ban-dep/KD-2026-0099")   # level 2 < min 4
    assert r.status_code == 404


def test_chua_co_ban_dep_404(tmp_path, monkeypatch):
    _users_file(tmp_path, monkeypatch)
    r = _dang_nhap("nv").get("/tai-ban-dep/KD-2026-0042")   # công khai nhưng chưa có PDF
    assert r.status_code == 404
    assert _dang_nhap("nv").get("/tai-ban-dep/XX-KHONG-CO").status_code == 404


def test_che_do_mo_khong_loc(monkeypatch):
    from claims_v2 import client_khach
    _chuan_bi_ban_dep()   # V2: claims thiếu bộ phận ≈ khách hệ cũ — không lọc quyền
    assert client_khach(app).get("/tai-ban-dep/KD-2026-0099").status_code == 200


def test_nang_cap_catalog_cu_15_cot():
    """Catalog thời trước Ý 4 (15 cột) → lần ghi kế tiếp tự thêm cột, dòng cũ đệm rỗng."""
    kho = Path(os.environ["KHO_TAI_LIEU"])
    kho.mkdir(parents=True, exist_ok=True)
    cu = kho / "_catalog.csv"
    with cu.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(CATALOG_HEADER[:15])                 # header cũ không có "Bản đẹp"
        w.writerow(["KD-1"] + ["x"] * 14)
    ghi_catalog(kho, ["KD-2"] + ["y"] * 14 + ["b_ban-dep.pdf"])
    dong = _catalog()
    assert dong[0] == CATALOG_HEADER
    # header giờ 18 cột (thêm Tầng nguồn + Tên nguồn — Supervisor); dòng cũ đệm rỗng tới 18
    assert len(dong[1]) == 18 and dong[1][15] == "" and dong[1][16] == "" and dong[1][17] == ""
    assert dong[2][15] == "b_ban-dep.pdf"


def test_nang_cap_catalog_14_cot_thoi_truoc_manh_b():
    """Bug bắt được khi KIỂM CHỨNG THẬT 19/07: catalog thật còn khuôn 14 cột (trước
    'Level tối thiểu') — phải CHÈN ô vị trí 8 để 'Tên file mới'/'Ngăn' về đúng chỗ,
    không đệm mù vào cuối."""
    kho = Path(os.environ["KHO_TAI_LIEU"])
    kho.mkdir(parents=True, exist_ok=True)
    with (kho / "_catalog.csv").open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(CATALOG_HEADER[:8] + CATALOG_HEADER[9:15])   # 14 cột, không min_level
        w.writerow(["KD-1", "2026-07-18", "t", "Kinh doanh", "Quy trình", "Còn hiệu lực",
                    "v1", "Mật", "mmo", "ai đó", "file.docx", "05_Kinh-doanh",
                    "false", "doc-id"])
    ghi_catalog(kho, ["KD-2"] + ["y"] * 15)
    dong = _catalog()
    assert dong[0] == CATALOG_HEADER
    assert dong[1][8] == ""                              # min_level chèn rỗng đúng vị trí
    assert dong[1][11] == "file.docx" and dong[1][12] == "05_Kinh-doanh"  # cột về đúng chỗ
    assert dong[1][13] == "false" and dong[1][14] == "doc-id" and dong[1][15] == ""


# ---- sửa bug treo: timeout cho lời gọi API ----

def test_llm_timeout_gan_vao_client(monkeypatch):
    from src.llm.openai_compatible import OpenAICompatibleProvider

    monkeypatch.setenv("LLM_TIMEOUT", "45")
    p = OpenAICompatibleProvider(model="m", api_key="k",
                                 base_url="http://localhost:9", mock=False)
    assert p.client.timeout == 45.0                      # không còn treo vô hạn theo SDK


def test_llm_retry_mac_dinh_0_khong_tu_thu_lai(monkeypatch):
    """Bẫy 06/08: SDK mặc định max_retries=2 LẶNG LẼ → flow phân tích nguồn 3 lời gọi
    × 3 lượt × 240s ≈ hơn 30 phút màn hình quay. Mặc định phải là 0 — lỗi nổi lên ngay."""
    from src.llm.openai_compatible import OpenAICompatibleProvider

    monkeypatch.delenv("LLM_RETRY", raising=False)
    p = OpenAICompatibleProvider(model="m", api_key="k",
                                 base_url="http://localhost:9", mock=False)
    assert p.client.max_retries == 0
    monkeypatch.setenv("LLM_RETRY", "1")                 # muốn SDK tự thử thêm → .env
    p2 = OpenAICompatibleProvider(model="m", api_key="k",
                                  base_url="http://localhost:9", mock=False)
    assert p2.client.max_retries == 1


# ---- sửa bug treo: bản đẹp chạy NỀN, /upload trả về ngay ----

def test_upload_tra_ve_ngay_ke_ca_ban_dep_hong(monkeypatch):
    """Nền ném lỗi cũng KHÔNG treo/hỏng request — message vẫn báo 'chế độ nền'."""
    import src.main as app_module

    def no_tung(*a, **k):
        raise RuntimeError("giả lập API nghẽn/treo")

    monkeypatch.setattr(app_module, "tao_ban_dep", no_tung)
    r = _upload("quy-trinh.docx", _docx_bytes())
    assert r.status_code == 200
    assert "chế độ nền" in r.json()["message"]
    assert _catalog()[1][15] == ""                       # nền hỏng → ô rỗng, chờ backfill


def test_upload_txt_khong_hua_ban_dep():
    r = _upload("ghi-chu.txt", "chữ".encode())
    assert r.status_code == 200
    assert "chế độ nền" not in r.json()["message"]       # không phải .docx — không hứa


def test_cap_nhat_ban_dep_dien_dung_dong():
    from src.main import cap_nhat_ban_dep_catalog

    kho = Path(os.environ["KHO_TAI_LIEU"])
    kho.mkdir(parents=True, exist_ok=True)
    ghi_catalog(kho, ["KD-1"] + [""] * 15)
    ghi_catalog(kho, ["KD-2"] + [""] * 15)
    cap_nhat_ban_dep_catalog(kho, "KD-2", "b_ban-dep.pdf")
    dong = _catalog()
    assert dong[1][15] == "" and dong[2][15] == "b_ban-dep.pdf"   # chỉ đúng dòng KD-2
